"""Build an Excel review workbook from reproducible CSV analysis outputs."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
OUTPUTS = ROOT / "p1-funnel-diagnostics" / "outputs"
DESTINATION = ROOT / "p1-funnel-diagnostics" / "merchant-onboarding-analysis.xlsx"

SHEETS = {
    "Funnel": OUTPUTS / "funnel_conversion.csv",
    "Activation Time": OUTPUTS / "activation_time.csv",
    "Retention": OUTPUTS / "cohort_retention.csv",
    "Channel Segment": OUTPUTS / "channel_segment_performance.csv",
}


def load_rows(path: Path) -> list[list[str]]:
    with path.open(encoding="utf-8") as handle:
        return list(csv.reader(handle))


def add_sheet(workbook: Workbook, title: str, rows: list[list[str]]) -> None:
    sheet = workbook.create_sheet(title)
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17212B")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = min(34, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        sheet.column_dimensions[column[0].column_letter].width = width
    table = Table(displayName=title.replace(" ", "") + "Table", ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    sheet.add_table(table)


def main() -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, path in SHEETS.items():
        add_sheet(workbook, title, load_rows(path))

    funnel = workbook["Funnel"]
    chart = BarChart()
    chart.type = "bar"
    chart.title = "Merchant onboarding funnel"
    chart.add_data(Reference(funnel, min_col=2, min_row=1, max_row=funnel.max_row), titles_from_data=True)
    chart.set_categories(Reference(funnel, min_col=1, min_row=2, max_row=funnel.max_row))
    chart.height, chart.width = 7, 13
    funnel.add_chart(chart, "F2")

    retention = workbook["Retention"]
    line = LineChart()
    line.title = "Retention by activation cohort"
    line.y_axis.title = "Retained %"
    line.x_axis.title = "Cohort observation"
    line.add_data(Reference(retention, min_col=5, min_row=1, max_row=retention.max_row), titles_from_data=True)
    line.height, line.width = 8, 15
    retention.add_chart(line, "G2")

    workbook.save(DESTINATION)
    print(f"Excel workbook written to {DESTINATION}")


if __name__ == "__main__":
    main()
